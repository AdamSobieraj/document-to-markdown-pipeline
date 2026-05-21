import io
import logging
import os
from typing import List, Union
import openpyxl
from langchain_core.documents import Document
from .base_parser import BaseDocumentParser

logger = logging.getLogger(__name__)


class XlsxParser(BaseDocumentParser):
    """Parser dla plików Excel, konwertuje arkusze na tabele Markdown."""

    def parse(self, file_source: Union[str, io.BytesIO, bytes], **kwargs) -> List[Document]:
        documents = []
        source_name = os.path.basename(file_source) if isinstance(file_source, str) else "strumień pamięci S3"

        try:
            # openpyxl natywnie wspiera zarówno ścieżki (str) jak i BytesIO
            wb = openpyxl.load_workbook(file_source, data_only=True)

            for i, sheet in enumerate(wb.worksheets):
                md_lines = [f"## Arkusz: {sheet.title}\n"]

                rows = list(sheet.iter_rows(values_only=True))
                if not rows:
                    continue

                for row_idx, row in enumerate(rows):
                    if any(cell is not None for cell in row):
                        # Czyszczenie i zamiana None na puste ciągi
                        clean_row = [str(cell).strip().replace('\n', ' ') if cell is not None else "" for cell in row]

                        # Wiersz tabeli w Markdown
                        md_row = "| " + " | ".join(clean_row) + " |"
                        md_lines.append(md_row)

                        # Separator nagłówka tabeli (tylko po pierwszym wierszu)
                        if row_idx == 0:
                            separator = "|-" + "-|-".join(
                                ["-" * len(c) if len(c) > 0 else "-" for c in clean_row]) + "-|"
                            md_lines.append(separator)

                if len(md_lines) > 1:
                    md_text = "\n".join(md_lines)
                    doc = Document(
                        page_content=md_text,
                        metadata={"page_number": i + 1, "sheet_name": sheet.title}
                    )
                    documents.append(doc)

            return documents
        except Exception as e:
            logger.error(f"Błąd przetwarzania XLSX ({source_name}): {e}")
            return []